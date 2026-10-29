# -*- coding: utf-8 -*-
"""
build_docx.py — 期末速通 Word 渲染器

把一份受控 Markdown（见 references/docx-spec.md）渲染成排版稳定的 .docx：
封面 / 可点击目录 / 章节书签 / 标签 / 提示框 / 表格 / 四角水印。

用法:
    python build_docx.py --src note.md --out note.docx
    python build_docx.py --src note.md --out note.docx --no-watermark
"""

import argparse
import re
import sys
import time
from copy import deepcopy
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

# ---------------------------------------------------------------- 设计常量

FONT_CN = "微软雅黑"
FONT_EN = "Calibri"
FONT_MONO = "Consolas"
INK = "262626"
INK_SOFT = "5A5A5A"
INK_FAINT = "8C8C8C"
RULE = "D8D8D8"
ACCENT = "1F4E5F"          # 默认强调色（深青），可由 front matter 覆盖
ACCENT_LIGHT = "EDF2F4"
PAGE_W_CM, PAGE_H_CM = 21.0, 29.7
SIZE_CM = 1.35
CORNER_CM = 0.4

TAG_STYLE = {
    "必考": ("FBE4DF", "A93226"),
    "重点": ("FBE4DF", "A93226"),
    "高频": ("FBEEDC", "9A6410"),
    "易错": ("F0EAF8", "64499B"),
    "秒杀": ("E4F3EA", "1E7350"),
    "中频": ("EAF0F5", "3C6180"),
    "了解": ("EFEFEF", "666666"),
    "不考": ("EFEFEF", "666666"),
}

CALL_TYPE = {
    "note":  ("F6F7F8", "8C969C", "提示"),
    "tip":   ("EEF6F2", "2E7D5B", "要点"),
    "key":   ("EEF6F2", "2E7D5B", "必记"),
    "warn":  ("FDF2EF", "C0503C", "易错"),
    "answer":("F2F5F9", "3B5F8A", "答案"),
    "summary":("FFF9EA", "C09A2E", "本章收口"),
    "overview":("F6F8F8", "1F4E5F", "本章速览"),
    "boss":  ("F6F8F8", "1F4E5F", "本章速览"),   # 兼容旧源文件，样式已改为普通卡片
    "grandma":("FFF7F2", "C9773F", "老奶版"),
    "muscle": ("F1F6F4", "2E7D5B", "薄肌版"),
}


# ---------------------------------------------------------------- 小工具

def _rgb(hexstr):
    return RGBColor.from_string(hexstr.replace("#", "").upper())


def _set_rfonts(rpr, ascii_font=None, east_font=None):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    if ascii_font:
        rf.set(qn("w:ascii"), ascii_font)
        rf.set(qn("w:hAnsi"), ascii_font)
        rf.set(qn("w:cs"), ascii_font)
    if east_font:
        rf.set(qn("w:eastAsia"), east_font)


def style_font(style, size=None, bold=None, color=None, ascii_font=FONT_EN,
               east_font=FONT_CN, italic=None):
    f = style.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = _rgb(color)
    rpr = style.element.get_or_add_rPr()
    _set_rfonts(rpr, ascii_font, east_font)


def shade(element, fill):
    """给 rPr / pPr / tcPr 加底纹。"""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def para_border(paragraph, edge="bottom", sz=6, color=RULE, space=4):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pbdr.append(el)


def add_field(paragraph, instr, placeholder="1", size=8.5, color=INK_FAINT):
    run = paragraph.add_run()
    run.font.size = Pt(size)
    run.font.color.rgb = _rgb(color)
    _set_rfonts(run._r.get_or_add_rPr(), FONT_EN, FONT_CN)
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = placeholder
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for el in (f1, it, f2, t, f3):
        run._r.append(el)


# ---------------------------------------------------------------- 内联解析

INLINE_RE = re.compile(
    r"(\*\*.+?\*\*|(?<!\*)\*(?!\s)[^*]+?(?<!\s)\*(?!\*)|`[^`]+?`"
    r"|\{\{[^{}]+?\}\}|\^\{[^{}]*\}|_\{[^{}]*\})"
)
TAG_RE = re.compile(r"\{\{([^{}]+)\}\}")


def strip_tags(text):
    return TAG_RE.sub("", text).strip()


def emit(paragraph, text, size=None, color=None, bold=False, italic=False):
    """写一段文字；文本里的 \\n 转成换行（用于选项、短句分行）。"""
    pieces = text.split("\n")
    for idx, piece in enumerate(pieces):
        if idx:
            paragraph.add_run().add_break()
        if not piece:
            continue
        run = paragraph.add_run(piece)
        if bold:
            run.bold = True
        if italic:
            run.italic = True
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = _rgb(color)


def render_inline(paragraph, text, size=None, color=None, bold=False, italic=False):
    parts = INLINE_RE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            emit(paragraph, part[2:-2], size, color, bold=True, italic=italic)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = paragraph.add_run(part[1:-1])
            r.font.name = FONT_MONO
            _set_rfonts(r._r.get_or_add_rPr(), FONT_MONO, FONT_CN)
            r.font.size = Pt((size or 10.5) - 1)
            r.font.color.rgb = _rgb("8A3A2E")
        elif TAG_RE.fullmatch(part):
            label = TAG_RE.fullmatch(part).group(1).strip()
            r = paragraph.add_run(f" {label} ")
            fill, fg = TAG_STYLE.get(label, ("EDEDED", "444444"))
            r.bold = True
            r.font.size = Pt((size or 10.5) - 1.5)
            r.font.color.rgb = _rgb(fg)
            _set_rfonts(r._r.get_or_add_rPr(), FONT_EN, FONT_CN)
            shade(r._r.get_or_add_rPr(), fill)
        elif part.startswith("^{") and part.endswith("}"):
            r = paragraph.add_run(part[2:-1])
            r.font.superscript = True
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = _rgb(color)
        elif part.startswith("_{") and part.endswith("}"):
            r = paragraph.add_run(part[2:-1])
            r.font.subscript = True
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = _rgb(color)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            emit(paragraph, part[1:-1], size, color, bold=bold, italic=True)
        else:
            emit(paragraph, part, size, color, bold=bold, italic=italic)


# ---------------------------------------------------------------- 文档解析

def parse_front_matter(text):
    meta, body = {}, text
    if text.startswith("---"):
        end = text.find("\n---", 3)
        if end != -1:
            block = text[3:end].strip("\n")
            for line in block.splitlines():
                if ":" in line:
                    k, v = line.split(":", 1)
                    meta[k.strip()] = v.strip().strip('"').strip("'")
            body = text[end + 4:]
    return meta, body.lstrip("\n")


def parse_blocks(lines):
    blocks, i = [], 0
    while i < len(lines):
        raw = lines[i]
        line = raw
        stripped = raw.strip()

        if not stripped:
            i += 1
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}", stripped):
            blocks.append(("hr",))
            i += 1
            continue

        # 图片
        m = re.match(r"^!\[(.*?)\]\((.+?)\)\s*$", stripped)
        if m:
            blocks.append(("img", m.group(1).strip(), m.group(2).strip()))
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            title = m.group(2).strip()
            summary = ""
            if "|" in title:
                title, summary = [x.strip() for x in title.split("|", 1)]
            blocks.append(("h", level, title, summary))
            i += 1
            continue

        # 提示框
        m = re.match(r"^:::\s*([a-zA-Z_]+)?\s*(.*)$", stripped)
        if m and not stripped.startswith("::::"):
            ctype = m.group(1) or "note"
            ctitle = m.group(2).strip()
            inner = []
            i += 1
            while i < len(lines) and lines[i].strip() != ":::":
                inner.append(lines[i])
                i += 1
            i += 1
            blocks.append(("callout", ctype, ctitle, parse_blocks(inner)))
            continue

        # 表格
        if stripped.startswith("|") and stripped.endswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                blocks.append(("table", rows))
            continue

        # 引用
        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip()[1:].strip())
                i += 1
            blocks.append(("quote", " ".join(buf)))
            continue

        # 列表
        m = re.match(r"^(\s*)([-*+]|\d+\.)\s+(.*)$", line)
        if m:
            indent = len(m.group(1).replace("\t", "  "))
            ordered = bool(re.match(r"\d+\.", m.group(2)))
            marker = m.group(2) if ordered else "-"
            blocks.append(("li", indent // 2, ordered, marker, m.group(3).strip()))
            i += 1
            continue

        # 选择题选项：A. B. C. D. 单独成段
        m = re.match(r"^([A-DＡ-Ｄ])[.、．]\s*(.*)$", stripped)
        if m:
            blocks.append(("option", m.group(1), m.group(2)))
            i += 1
            continue

        # 普通段落（合并软换行；行尾两个空格 = 强制换行）
        buf = [line.rstrip()]
        hard = [line.endswith("  ")]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            is_special = (
                not nxt
                or nxt.startswith(("#", "|", ">", ":::"))
                or re.fullmatch(r"-{3,}", nxt) is not None
                or re.match(r"^(\s*)([-*+]|\d+\.)\s+", lines[i]) is not None
                or re.match(r"^[A-DＡ-Ｄ][.、．]\s*", nxt) is not None
            )
            if is_special:
                break
            buf.append(lines[i].strip())
            hard.append(lines[i].rstrip("\n").endswith("  "))
            i += 1
        text = ""
        for idx, seg in enumerate(buf):
            if idx and hard[idx - 1]:
                text += "\n"
            elif text and re.search(r"[A-Za-z0-9,.;:!?]$", text) and re.match(r"^[A-Za-z0-9(]", seg):
                text += " "
            text += seg
        blocks.append(("p", text))
    return blocks


# ---------------------------------------------------------------- 渲染器

class Builder:
    def __init__(self, meta, watermark=None, base_dir="."):
        self.meta = meta
        self.accent = meta.get("accent", ACCENT)
        self.doc = Document()
        self.bookmark_seq = 0
        self.headings = []       # (level, title, anchor)
        self.watermark = watermark
        self.base_dir = Path(base_dir)
        self.toc_depth = int(meta.get("toc_depth", "1"))
        self._setup_page()
        self._setup_styles()

    # -- 页面与样式 -------------------------------------------------
    def _setup_page(self):
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(PAGE_W_CM), Cm(PAGE_H_CM)
        s.top_margin = Cm(2.15)
        s.bottom_margin = Cm(2.15)
        s.left_margin = Cm(2.4)
        s.right_margin = Cm(2.4)
        s.header_distance = Cm(0.6)
        s.footer_distance = Cm(0.75)

    def _setup_styles(self):
        st = self.doc.styles
        normal = st["Normal"]
        style_font(normal, size=10.5, color=INK)
        normal.paragraph_format.line_spacing = 1.55
        normal.paragraph_format.space_after = Pt(3)
        normal.paragraph_format.space_before = Pt(0)

        def mk(name, base="Normal"):
            try:
                return st[name]
            except KeyError:
                from docx.enum.style import WD_STYLE_TYPE
                return st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)

        h1 = mk("FC H1")
        h1.base_style = st["Normal"]
        style_font(h1, size=17, bold=True, color="1A1A1A")
        no_break = str(self.meta.get("page_break", "true")).lower() == "false"
        h1.paragraph_format.space_before = Pt(20 if no_break else 0)
        h1.paragraph_format.space_after = Pt(6)
        h1.paragraph_format.page_break_before = not no_break
        h1.paragraph_format.keep_with_next = True
        h1.paragraph_format.line_spacing = 1.2

        h2 = mk("FC H2")
        h2.base_style = st["Normal"]
        style_font(h2, size=13, bold=True, color=self.accent)
        h2.paragraph_format.space_before = Pt(14)
        h2.paragraph_format.space_after = Pt(4)
        h2.paragraph_format.keep_with_next = True
        h2.paragraph_format.line_spacing = 1.25

        h3 = mk("FC H3")
        h3.base_style = st["Normal"]
        style_font(h3, size=11, bold=True, color="1A1A1A")
        h3.paragraph_format.space_before = Pt(9)
        h3.paragraph_format.space_after = Pt(2)
        h3.paragraph_format.keep_with_next = True
        h3.paragraph_format.line_spacing = 1.3

        h4 = mk("FC H4")
        h4.base_style = st["Normal"]
        style_font(h4, size=10.5, bold=True, color=INK_SOFT)
        h4.paragraph_format.space_before = Pt(6)
        h4.paragraph_format.space_after = Pt(2)
        h4.paragraph_format.keep_with_next = True

        for name, lvl in (("FC H1", 0), ("FC H2", 1), ("FC H3", 2), ("FC H4", 3)):
            ppr = st[name].element.get_or_add_pPr()
            ol = OxmlElement("w:outlineLvl")
            ol.set(qn("w:val"), str(lvl))
            ppr.append(ol)

    # -- 封面 / 目录 -------------------------------------------------
    def cover(self):
        for _ in range(3):
            self.doc.add_paragraph()

        kicker = self.doc.add_paragraph()
        kicker.paragraph_format.space_after = Pt(10)
        r = kicker.add_run((self.meta.get("kicker", "FINAL SPRINT · 期末速通")).upper())
        r.font.size = Pt(8.5)
        r.font.color.rgb = _rgb(INK_FAINT)
        r.bold = True
        rpr = r._r.get_or_add_rPr()
        _set_rfonts(rpr, FONT_EN, FONT_CN)
        sp = OxmlElement("w:spacing"); sp.set(qn("w:val"), "80"); rpr.append(sp)

        title = self.doc.add_paragraph()
        title.paragraph_format.space_after = Pt(6)
        title.paragraph_format.line_spacing = 1.15
        rt = title.add_run(self.meta.get("title", "期末速通"))
        rt.font.size = Pt(30)
        rt.bold = True
        rt.font.color.rgb = _rgb("1A1A1A")
        _set_rfonts(rt._r.get_or_add_rPr(), FONT_EN, FONT_CN)

        subtitle = self.meta.get("subtitle", "")
        if subtitle:
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(18)
            r = p.add_run(subtitle)
            r.font.size = Pt(12)
            r.font.color.rgb = _rgb(INK_SOFT)
            _set_rfonts(r._r.get_or_add_rPr(), FONT_EN, FONT_CN)

        rule = self.doc.add_paragraph()
        rule.paragraph_format.space_after = Pt(16)
        para_border(rule, "bottom", sz=14, color=self.accent, space=1)

        meta_rows = [
            ("课程", self.meta.get("course", "")),
            ("版本", self.meta.get("track", "")),
            ("适用", self.meta.get("audience", "")),
            ("生成日期", self.meta.get("date", "")),
        ]
        meta_rows = [(k, v) for k, v in meta_rows if v]
        if meta_rows:
            table = self.doc.add_table(rows=0, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            table.autofit = False
            table.columns[0].width = Cm(3.0)
            table.columns[1].width = Cm(10.6)
            for k, v in meta_rows:
                row = table.add_row()
                c0, c1 = row.cells
                c0.width, c1.width = Cm(3.0), Cm(10.6)
                p0 = c0.paragraphs[0]
                r0 = p0.add_run(k)
                r0.font.size = Pt(9.5)
                r0.font.color.rgb = _rgb(INK_FAINT)
                p1 = c1.paragraphs[0]
                r1 = p1.add_run(v)
                r1.font.size = Pt(10)
                r1.font.color.rgb = _rgb(INK)
                for p in (p0, p1):
                    p.paragraph_format.space_after = Pt(4)

        for _ in range(4):
            self.doc.add_paragraph()

        note = self.doc.add_paragraph()
        r = note.add_run(self.meta.get(
            "cover_note", "以题目为准，以得分为目标。看不懂的地方，先背结论，再补原理。"))
        r.font.size = Pt(9)
        r.font.color.rgb = _rgb(INK_FAINT)
        _set_rfonts(r._r.get_or_add_rPr(), FONT_EN, FONT_CN)

    def toc(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        bm = OxmlElement("w:bookmarkStart")
        bm.set(qn("w:id"), "9000")
        bm.set(qn("w:name"), "_FC_TOC_TOP")
        p._p.append(bm)
        r = p.add_run("目录")
        r.font.size = Pt(17); r.bold = True
        r.font.color.rgb = _rgb("1A1A1A")
        _set_rfonts(r._r.get_or_add_rPr(), FONT_EN, FONT_CN)
        bme = OxmlElement("w:bookmarkEnd"); bme.set(qn("w:id"), "9000")
        p._p.append(bme)
        para_border(p, "bottom", sz=8, color=RULE, space=4)

        hint = self.doc.add_paragraph()
        hint.paragraph_format.space_before = Pt(4)
        hint.paragraph_format.space_after = Pt(12)
        rh = hint.add_run("点击任意章节标题可直接跳转 · 每章末尾可返回目录")
        rh.font.size = Pt(8.5); rh.font.color.rgb = _rgb(INK_FAINT)
        _set_rfonts(rh._r.get_or_add_rPr(), FONT_EN, FONT_CN)

        for level, title, summary, anchor in self.headings:
            if level > self.toc_depth:
                continue
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.space_before = Pt(7 if level == 1 else 2)
            pf.space_after = Pt(1 if level == 1 else 0)
            pf.left_indent = Cm(0 if level == 1 else 0.85)
            pf.line_spacing = 1.3
            add_internal_link(para, strip_tags(title), anchor,
                              color="1A1A1A" if level == 1 else INK_SOFT,
                              size=10.5 if level == 1 else 9.5,
                              bold=(level == 1))
            if level == 1 and summary:
                sp = self.doc.add_paragraph()
                sp.paragraph_format.left_indent = Cm(0)
                sp.paragraph_format.space_before = Pt(0)
                sp.paragraph_format.space_after = Pt(6)
                sr = sp.add_run(strip_tags(summary))
                sr.font.size = Pt(8.5)
                sr.font.color.rgb = _rgb(INK_FAINT)
                _set_rfonts(sr._r.get_or_add_rPr(), FONT_EN, FONT_CN)

    # -- 正文 -------------------------------------------------------
    def heading(self, level, title, summary, anchor):
        self.bookmark_seq += 1
        p = self.doc.add_paragraph(style=f"FC H{level}")
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self.bookmark_seq))
        start.set(qn("w:name"), anchor)
        p._p.append(start)

        render_inline(p, title, size={1: 17, 2: 13, 3: 11, 4: 10.5}[level],
                      color=("1A1A1A" if level in (1, 3, 4) else self.accent),
                      bold=True)

        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(self.bookmark_seq))
        p._p.append(end)

        if level == 1:
            para_border(p, "bottom", sz=10, color=self.accent, space=6)
        if summary:
            sp = self.doc.add_paragraph()
            sp.paragraph_format.space_before = Pt(2)
            sp.paragraph_format.space_after = Pt(8)
            r = sp.add_run(strip_tags(summary))
            r.font.size = Pt(9)
            r.font.color.rgb = _rgb(INK_FAINT)
            r.italic = True
            _set_rfonts(r._r.get_or_add_rPr(), FONT_EN, FONT_CN)
        return anchor

    def paragraph(self, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        render_inline(p, text)
        return p

    def quote(self, text):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.45)
        pf.space_before = Pt(3)
        pf.space_after = Pt(5)
        para_border(p, "left", sz=18, color="C9D2D6", space=8)
        r = p.add_run("")
        render_inline(p, text, color=INK_SOFT)
        return p

    def option(self, letter, text):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.85)
        pf.first_line_indent = Cm(-0.45)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.4
        r = p.add_run(f"{letter}. ")
        r.bold = True
        r.font.color.rgb = _rgb(INK_SOFT)
        render_inline(p, text)
        return p

    def list_item(self, level, ordered, marker, text):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.62 + 0.55 * level)
        pf.first_line_indent = Cm(-0.42)
        pf.space_before = Pt(1)
        pf.space_after = Pt(1)
        pf.line_spacing = 1.4
        if ordered:
            label = marker
        else:
            label = "•" if level == 0 else "–"
        r = p.add_run(label + "\t")
        r.font.size = Pt(10.5)
        r.font.color.rgb = _rgb(self.accent if level == 0 else INK_FAINT)
        r.bold = ordered and level == 0
        render_inline(p, text)

    def hr(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        para_border(p, "bottom", sz=4, color=RULE, space=1)

    def image(self, caption, src):
        path = Path(src)
        if not path.is_absolute():
            path = self.base_dir / path
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = bool(caption)

        if not path.exists():
            r = p.add_run(f"[图缺失: {src}]")
            r.font.size = Pt(9)
            r.font.color.rgb = _rgb("B03A2E")
        else:
            try:
                from docx.image.image import Image as DocxImage
                info = DocxImage.from_file(str(path))
                dpi = info.horz_dpi or 96
                width_cm = info.width / dpi * 2.54
                height_cm = info.height / dpi * 2.54
                # 同时限制宽和高，避免竖长图撑出页面。
                # 只传 width，高度按原比例走；因此不能再把 width 抬回最小值，
                # 否则超长竖图会重新撑出页面。
                scale = min(1.0, 15.4 / max(width_cm, 0.1), 17.0 / max(height_cm, 0.1))
                width_cm = max(width_cm * scale, 0.1)
                p.add_run().add_picture(str(path), width=Cm(width_cm))
            except Exception:
                r = p.add_run(f"[图片无法读取: {src}]")
                r.font.size = Pt(9)
                r.font.color.rgb = _rgb("B03A2E")

        if caption:
            cp = self.doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_before = Pt(0)
            cp.paragraph_format.space_after = Pt(9)
            r = cp.add_run(caption)
            r.font.size = Pt(8.5)
            r.font.color.rgb = _rgb(INK_FAINT)
            _set_rfonts(r._r.get_or_add_rPr(), FONT_EN, FONT_CN)
        else:
            p.paragraph_format.space_after = Pt(8)

    def callout(self, ctype, title, inner_blocks):
        fill, edge, default_title = CALL_TYPE.get(ctype, CALL_TYPE["note"])
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        tbl = table._tbl
        tblPr = tbl.tblPr

        # 宽度、边框、底纹
        tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), "5000"); tw.set(qn("w:type"), "pct")
        tblPr.append(tw)
        borders = OxmlElement("w:tblBorders")
        for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge_name}")
            el.set(qn("w:val"), "single" if edge_name == "left" else "nil")
            el.set(qn("w:sz"), "24" if edge_name == "left" else "0")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), edge)
            borders.append(el)
        tblPr.append(borders)

        cell = table.cell(0, 0)
        cell.width = Cm(PAGE_W_CM - 4.8)
        tcPr = cell._tc.get_or_add_tcPr()
        shade(tcPr, fill)
        mar = OxmlElement("w:tcMar")
        for side, val in (("top", 110), ("left", 200), ("bottom", 110), ("right", 160)):
            el = OxmlElement(f"w:{side}"); el.set(qn("w:w"), str(val)); el.set(qn("w:type"), "dxa")
            mar.append(el)
        tcPr.append(mar)

        # 提示框允许跨页拆分，避免整块下移造成半页空白

        first = cell.paragraphs[0]
        head = title or default_title
        if head:
            hp = first
            hp.paragraph_format.space_after = Pt(3)
            r = hp.add_run(head)
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = _rgb(edge)
            _set_rfonts(r._r.get_or_add_rPr(), FONT_EN, FONT_CN)

        if inner_blocks:
            for idx, block in enumerate(inner_blocks):
                if block[0] == "li":
                    p = cell.add_paragraph()
                    _, lvl, ordered, marker, text = block
                    pf = p.paragraph_format
                    pf.left_indent = Cm(0.5)
                    pf.first_line_indent = Cm(-0.4)
                    pf.space_after = Pt(1)
                    rr = p.add_run(("•" if not ordered else marker) + " ")
                    rr.font.color.rgb = _rgb(edge)
                    render_inline(p, text)
                elif block[0] == "option":
                    p = cell.add_paragraph()
                    p.paragraph_format.left_indent = Cm(0.55)
                    p.paragraph_format.first_line_indent = Cm(-0.45)
                    rr = p.add_run(f"{block[1]}. ")
                    rr.bold = True
                    render_inline(p, block[2])
                elif block[0] == "p":
                    p = cell.add_paragraph()
                    p.paragraph_format.space_after = Pt(2)
                    render_inline(p, block[1])
                elif block[0] == "quote":
                    p = cell.add_paragraph()
                    render_inline(p, block[1], color=INK_SOFT)
                elif block[0] == "h":
                    p = cell.add_paragraph()
                    render_inline(p, block[2], size=10.5, bold=True, color=edge)
                elif block[0] == "hr":
                    continue

        if ctype in ("overview", "boss"):
            for cp in cell.paragraphs:
                cp.paragraph_format.line_spacing = 1.35

        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)
        spacer.paragraph_format.line_spacing = 0.6

    def table(self, rows):
        ncols = max(len(r) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.allow_autofit = False
        tblPr = table._tbl.tblPr
        tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), "5000"); tw.set(qn("w:type"), "pct")
        tblPr.append(tw)
        layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)

        # 按内容长度分配列宽，避免长文本被挤成竖条
        def text_w(s):
            s = strip_tags(s)
            return sum(2 if ord(ch) > 127 else 1 for ch in s)

        usable = PAGE_W_CM - 4.8
        raw = [max(text_w(r[i]) if i < len(r) else 0 for r in rows) for i in range(ncols)]
        weights = [min(max(v, 6), 44) for v in raw]
        widths = [usable * v / sum(weights) for v in weights]
        widths = [max(2.2, x) for x in widths]
        scale = usable / sum(widths)
        widths = [x * scale for x in widths]
        for i, cm in enumerate(widths):
            table.columns[i].width = Cm(cm)
        for row in table.rows:
            for i, cm in enumerate(widths):
                row.cells[i].width = Cm(cm)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            if edge in ("top", "bottom", "insideH"):
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "4" if edge == "insideH" else "6")
                el.set(qn("w:color"), RULE if edge == "insideH" else "BFBFBF")
            else:
                el.set(qn("w:val"), "nil")
            el.set(qn("w:space"), "0")
            borders.append(el)
        tblPr.append(borders)

        for r_idx, row in enumerate(rows):
            for c_idx in range(ncols):
                cell = table.cell(r_idx, c_idx)
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                text = row[c_idx] if c_idx < len(row) else ""
                if r_idx == 0:
                    tcPr = cell._tc.get_or_add_tcPr()
                    shade(tcPr, "EEF1F2")
                    render_inline(p, text, size=9.5, bold=True)
                else:
                    render_inline(p, text, size=9.5)
        # 表头跨页重复 + 行内不断页 + 小表整体不拆
        trPr = table.rows[0]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:tblHeader"))
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if len(rows) <= 20:
            for row in table.rows[:-1]:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True

        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(3)
        spacer.paragraph_format.line_spacing = 0.6

    def back_to_toc(self):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(8)
        add_internal_link(p, "↑ 返回目录", "_FC_TOC_TOP", color=INK_FAINT, size=8)

    def build(self, blocks):
        self.cover()

        # 预扫描标题，生成目录（锚点必须与正文完全一致）
        self.headings = []
        seq = 0
        for b in blocks:
            if b[0] == "h":
                seq += 1
                self.headings.append((b[1], b[2], b[3], f"_FC_H{b[1]}_{seq}"))

        self.toc()

        hidx = 0
        chapter_open = False
        for b in blocks:
            kind = b[0]
            if kind == "h":
                if chapter_open and b[1] == 1:
                    self.back_to_toc()
                anchor = self.headings[hidx][3]
                hidx += 1
                self.heading(b[1], b[2], b[3], anchor)
                if b[1] == 1:
                    chapter_open = True
            elif kind == "p":
                self.paragraph(b[1])
            elif kind == "li":
                self.list_item(*b[1:])
            elif kind == "option":
                self.option(*b[1:])
            elif kind == "quote":
                self.quote(b[1])
            elif kind == "table":
                self.table(b[1])
            elif kind == "img":
                self.image(b[1], b[2])
            elif kind == "callout":
                self.callout(*b[1:])
            elif kind == "hr":
                self.hr()
        # 末章不再加返回链接，避免单独占一页

    # -- 页眉页脚 ---------------------------------------------------
    def decorations(self, watermark=None):
        s = self.doc.sections[0]
        if watermark and Path(watermark).exists():
            self._corners(s, watermark)
        footer = s.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.space_before = Pt(0)
        add_field(footer, "PAGE", "1", size=8.5)
        r = footer.add_run("  /  ")
        r.font.size = Pt(8.5); r.font.color.rgb = _rgb(INK_FAINT)
        add_field(footer, "NUMPAGES", "1", size=8.5)
        if self.meta.get("footer"):
            r2 = footer.add_run("　·　" + self.meta["footer"])
            r2.font.size = Pt(8.5); r2.font.color.rgb = _rgb(INK_FAINT)
            _set_rfonts(r2._r.get_or_add_rPr(), FONT_EN, FONT_CN)

    def _corners(self, section, image):
        header = section.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        # 页眉段落压到 1pt，避免把正文往下推
        p.paragraph_format.line_spacing = Pt(1)

        img = Cm(SIZE_CM)
        x_left = Cm(CORNER_CM)
        x_right = Cm(PAGE_W_CM - CORNER_CM - SIZE_CM)
        y_top = Cm(CORNER_CM)
        y_bottom = Cm(PAGE_H_CM - CORNER_CM - SIZE_CM)

        for x, y in ((x_left, y_top), (x_right, y_top), (x_left, y_bottom), (x_right, y_bottom)):
            run = p.add_run()
            run.add_picture(str(image), width=img)
            drawing = run._r.find(qn("w:drawing"))
            inline = drawing[0]
            anchor = self._anchor_from_inline(inline, int(x), int(y))
            drawing.replace(inline, anchor)

    @staticmethod
    def _anchor_from_inline(inline, x_emu, y_emu):
        anchor = OxmlElement("wp:anchor")
        for k, v in dict(distT="0", distB="0", distL="0", distR="0", simplePos="0",
                         relativeHeight="1", behindDoc="1", locked="1",
                         layoutInCell="1", allowOverlap="1").items():
            anchor.set(k, v)
        simple = OxmlElement("wp:simplePos"); simple.set("x", "0"); simple.set("y", "0")
        anchor.append(simple)
        ph = OxmlElement("wp:positionH"); ph.set("relativeFrom", "page")
        off = OxmlElement("wp:posOffset"); off.text = str(x_emu); ph.append(off)
        anchor.append(ph)
        pv = OxmlElement("wp:positionV"); pv.set("relativeFrom", "page")
        off = OxmlElement("wp:posOffset"); off.text = str(y_emu); pv.append(off)
        anchor.append(pv)
        anchor.append(deepcopy(inline.find(qn("wp:extent"))))
        eff = inline.find(qn("wp:effectExtent"))
        if eff is None:
            eff = OxmlElement("wp:effectExtent")
            for k in ("l", "t", "r", "b"):
                eff.set(k, "0")
        anchor.append(deepcopy(eff))
        anchor.append(OxmlElement("wp:wrapNone"))
        anchor.append(deepcopy(inline.find(qn("wp:docPr"))))
        cnv = inline.find(qn("wp:cNvGraphicFramePr"))
        if cnv is not None:
            anchor.append(deepcopy(cnv))
        anchor.append(deepcopy(inline.find(qn("a:graphic"))))
        return anchor


def add_internal_link(paragraph, text, anchor, color="1A1A1A", size=10.5, bold=False):
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("w:anchor"), anchor)
    hl.set(qn("w:history"), "1")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    _set_rfonts(rpr, FONT_EN, FONT_CN)
    if bold:
        rpr.append(OxmlElement("w:b"))
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rpr.append(sz)
    szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), str(int(size * 2))); rpr.append(szcs)
    c = OxmlElement("w:color"); c.set(qn("w:val"), color); rpr.append(c)
    r.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
    r.append(t)
    hl.append(r)
    paragraph._p.append(hl)


# ---------------------------------------------------------------- 入口

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--watermark", default=str(Path(__file__).resolve().parent.parent / "assets" / "mascot_faded.png"))
    ap.add_argument("--no-watermark", action="store_true")
    args = ap.parse_args()

    raw = Path(args.src).read_text(encoding="utf-8")
    meta, body = parse_front_matter(raw)
    blocks = parse_blocks(body.splitlines())

    b = Builder(meta, base_dir=Path(args.src).resolve().parent)
    b.build(blocks)
    b.decorations(None if args.no_watermark else args.watermark)
    b.doc.core_properties.title = meta.get("title", "")
    b.doc.core_properties.author = meta.get("author", "")

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)

    saved = out
    for attempt in range(3):
        try:
            b.doc.save(str(saved))
            break
        except PermissionError:
            if attempt == 2:
                saved = out.with_name(f"{out.stem}_new{out.suffix}")
                b.doc.save(str(saved))
                print(f"WARN 目标文件被占用（网盘同步/预览），已另存为: {saved}")
            else:
                time.sleep(1.5)

    print(f"OK {saved}  ({saved.stat().st_size/1024:.0f} KB)")


if __name__ == "__main__":
    main()
